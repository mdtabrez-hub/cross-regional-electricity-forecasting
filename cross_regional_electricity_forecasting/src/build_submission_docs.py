"""Generate submission documents (cover letter, highlights, declarations, reviewer-response
template) as editable Word files, filling numbers from results/manuscript_numbers.json.

Usage: python -m src.build_submission_docs
"""
from __future__ import annotations

import json

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .common import RESULTS, ROOT

SUB = ROOT / "submission"
MS = ROOT / "manuscript"
N = json.load(open(RESULTS / "manuscript_numbers.json"))
TITLE = ("Forecasting Electricity Demand in Unseen Regions: A Leave-Regions-Out Benchmark of Global Models, "
         "Short-Window Adaptation and Conformal Calibration across 48 U.S. Balancing Authorities")


def doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")
    for s in ("Heading 1", "Heading 2", "Title"):
        d.styles[s].font.name = "Times New Roman"
        d.styles[s].font.color.rgb = RGBColor(0, 0, 0)
    return d


def para(d, text, bold=False, italic=False, align=None, size=None):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def cover_letter():
    d = doc()
    para(d, "[AUTHOR ACTION: date]")
    para(d, "The Editor-in-Chief\nInternational Journal of Forecasting [or the selected journal]")
    para(d, f"Dear Editor,")
    para(d, f"We submit the manuscript \"{TITLE}\" for consideration as a research article.")
    para(d, (
        "Forecasting models are increasingly expected to serve power systems they were not trained on: new or "
        "reorganised balancing areas, utilities joining markets, and small balancing authorities without forecasting "
        "staff. The manuscript reports a leakage-resistant benchmark, with a protocol frozen in version control before test evaluation, on public EIA Form 930 data for "
        f"{N['data.n_regions']} U.S. balancing authorities (2019–2025). It separates geographic from temporal "
        "distribution shift for both point accuracy and prediction-interval calibration."))
    para(d, "The main findings are:")
    for t in [
        f"a pooled, instance-normalised LightGBM model loses a median {N['test.h24.nMAE.LGB_G_LRO.LGB_G_IN.median_rel_change']}% "
        f"(day-ahead) to {N['test.h1.nMAE.LGB_G_LRO.LGB_G_IN.median_rel_change']}% (hour-ahead) relative accuracy on balancing "
        f"authorities it has never seen ({N['revsptest.h24.LGB_G_LRS.LGB_G_IN.median_rel_change']}% day-ahead when whole regions "
        "are held out), matching local models day-ahead and beating them intraday, with a tail of larger losses intraday;",
        "fine-tuning on one to thirteen weeks of target data degraded accuracy at every onboarding date tested;",
        "static conformal intervals are near-nominal on average but under-cover in a fifth to a third of systems, "
        "whether or not the system was seen in training; and",
        "adaptive conformal inference restores per-system coverage, but its practical value depends on a step size "
        "suited to the feedback delay; choosing it on other systems worked well.",
    ]:
        d.add_paragraph(t, style="List Bullet")
    para(d, (
        "The contribution is an evaluation framework and empirical evidence rather than a new algorithm, which we "
        "believe fits the journal's interest in forecast evaluation, probabilistic calibration and reproducible energy "
        "forecasting. All code, the frozen protocol, a decision log of every post-freeze amendment, predictions and "
        "metrics are openly available [AUTHOR ACTION: repository/DOI]."))
    para(d, (
        "The manuscript is original, has not been published previously and is not under consideration elsewhere. "
        "All authors approved the submission. [AUTHOR ACTION: confirm; state any preprint posting; declare competing "
        "interests and generative-AI assistance as required by the journal.]"))
    para(d, "[AUTHOR ACTION: optional suggested reviewers with institutional e-mails and no conflicts of interest.]")
    para(d, "Yours sincerely,\n[Corresponding author name, affiliation, e-mail, ORCID]")
    d.save(SUB / "cover_letter.docx")


def highlights():
    g24, g1 = N["test.h24.nMAE.LGB_G_LRO.LGB_G_IN.median_rel_change"], N["test.h1.nMAE.LGB_G_LRO.LGB_G_IN.median_rel_change"]
    sp = N["revsptest.h24.LGB_G_LRS.LGB_G_IN.median_rel_change"]
    u24 = N["prob.h24.unseen-region.90.CQR.all.share_regions_undercover_5pp"]
    u1 = N["prob.h1.unseen-region.90.CQR.all.share_regions_undercover_5pp"]
    hl = [
        "Leakage-resistant benchmark: forecasting 48 U.S. balancing authorities never seen",
        f"Unseen systems cost {g24}–{g1}% median accuracy; more if whole regions are new",
        "Fine-tuning on 1–13 weeks of target data degraded accuracy at every onboarding date",
        f"Static conformal intervals under-covered in {round(float(u24))}–{round(float(u1))}% of systems, seen or unseen",
        "Online conformal with a step size tuned on other systems restored coverage",
    ]
    d = doc()
    para(d, "Highlights", bold=True)
    for h in hl:
        assert len(h) <= 85, (h, len(h))
        d.add_paragraph(h, style="List Bullet")
    para(d, "(Each highlight ≤ 85 characters including spaces, as required by Applied Energy; IJF does not require "
            "highlights — verify against the selected journal's Guide for Authors.)", italic=True, size=10)
    d.save(SUB / "highlights.docx")
    (SUB / "highlights.txt").write_text("\n".join(hl) + "\n")


def declarations():
    d = doc()
    para(d, "Declarations template (to be completed by the authors)", bold=True, size=14)
    for head, body in [
        ("Data availability statement", (MS / "sections" / "08_backmatter.md").read_text().split("# Code availability")[0]
         .replace("# Data availability statement {.unnumbered}", "").strip()),
        ("Code availability statement", "All code to reproduce the study is available at [AUTHOR ACTION: repository URL/DOI]; "
         "`python -m src.run_all` reproduces every number, table and figure."),
        ("Author contributions (CRediT)", "[Author 1]: Conceptualization; Methodology; Software; Formal analysis; Writing – "
         "original draft. [Author 2]: Supervision; Writing – review & editing. [AUTHOR ACTION: edit to reflect actual roles.]"),
        ("Funding", "[AUTHOR ACTION] e.g. \"This research did not receive any specific grant from funding agencies in the "
         "public, commercial, or not-for-profit sectors.\""),
        ("Declaration of competing interest", "[AUTHOR ACTION] e.g. \"The authors declare that they have no known competing "
         "financial interests or personal relationships that could have appeared to influence the work reported in this paper.\""),
        ("Declaration of generative AI use", "[AUTHOR ACTION] The manuscript, code and analyses were prepared with substantial "
         "assistance from an AI system (Claude, Anthropic). Complete the journal's required AI-use statement; authors must "
         "verify all content and take full responsibility."),
        ("Ethics", "The study uses public, aggregate electricity-system data only; no human participants or personal data."),
    ]:
        d.add_heading(head, level=2)
        para(d, body)
    d.save(SUB / "declarations_template.docx")


def response_template():
    rc = MS / "revision_matrix.csv"
    d = doc()
    para(d, "Response to Reviewers — PREPARATION EXERCISE", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d, ("IMPORTANT: The comments below come from an INTERNAL, SIMULATED review produced during manuscript preparation. "
             "They are NOT comments from any journal or real referee. This document is a template and rehearsal for "
             "responding to genuine reviews after submission."), bold=True)
    if rc.exists():
        m = pd.read_csv(rc)
        for _, r in m.iterrows():
            d.add_heading(f"Comment {r['id']} ({r['severity']})", level=2)
            para(d, f"Comment: {r['reviewer_comment']}", italic=True)
            para(d, f"Response: {r['author_response']}")
            para(d, f"Action taken: {r['action_taken']} (section: {r['modified_section']}; verification: {r['verification_status']})")
    else:
        para(d, "[Revision matrix not yet available]")
    d.save(MS / "response_to_reviewers_template.docx")


def main():
    SUB.mkdir(exist_ok=True)
    cover_letter()
    highlights()
    declarations()
    response_template()
    print("submission documents written")


if __name__ == "__main__":
    main()
